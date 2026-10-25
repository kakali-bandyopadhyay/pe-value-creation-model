from __future__ import annotations

"""Recalculate scenarios in LibreOffice and independently verify workbook returns."""

import argparse
import math
import shutil
import subprocess
import sys
import uuid
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from build.build_model import build_workbook

WORKBOOK = ROOT / "model" / "PE_Value_Creation_Model_synthetic.xlsx"
SCENARIOS = ("Base", "Downside", "Upside")
INIT_ROWS = {
    "Pricing and mix": {"Base": 7, "Downside": 8, "Upside": 9},
    "Commercial efficiency": {"Base": 14, "Downside": 15, "Upside": 16},
    "AI finance automation": {"Base": 21, "Downside": 22, "Upside": 23},
}
OPERATING_ROWS = {"Pricing and mix": 17, "Commercial efficiency": 23, "AI finance automation": 32}


def office_executable() -> str:
    found = shutil.which("soffice.com")
    if found:
        return found
    candidate = Path(r"C:\Program Files\LibreOffice\program\soffice.com")
    if candidate.exists():
        return str(candidate)
    found = shutil.which("soffice")
    if found:
        return found
    raise FileNotFoundError("LibreOffice soffice.com was not found. Install LibreOffice or add soffice to PATH.")


def recalculate(source: Path, output_dir: Path, profile_dir: Path) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    profile_dir.mkdir(parents=True, exist_ok=True)
    profile_uri = "file:///" + profile_dir.as_posix()
    cmd = [office_executable(), f"-env:UserInstallation={profile_uri}", "--headless", "--convert-to", "xlsx", "--outdir", str(output_dir), str(source)]
    result = subprocess.run(cmd, capture_output=True, text=True, timeout=120)
    target = output_dir / source.name
    if result.returncode != 0 or not target.exists():
        raise RuntimeError(f"LibreOffice recalculation failed. Exit={result.returncode}\n{result.stdout}\n{result.stderr}")
    return target


def fresh_run_dir(name: str) -> Path:
    parent = ROOT / "out"
    parent.mkdir(parents=True, exist_ok=True)
    run_dir = parent / f"{name}_{uuid.uuid4().hex[:8]}"
    run_dir.mkdir()
    return run_dir


def xirr(cashflows: list[float], dates: list[date | datetime]) -> float:
    dated = [(float(value), dt.date() if isinstance(dt, datetime) else dt) for value, dt in zip(cashflows, dates) if value]
    if not dated or not any(v < 0 for v, _ in dated) or not any(v > 0 for v, _ in dated):
        raise ValueError("XIRR requires at least one negative and one positive cash flow.")
    start = dated[0][1]
    def npv(rate: float) -> float:
        return sum(value / ((1 + rate) ** ((dt - start).days / 365.0)) for value, dt in dated)
    lo, hi = -0.999999, 100.0
    if npv(lo) * npv(hi) > 0:
        raise ValueError("Could not bracket an XIRR root for the dated cash flows.")
    for _ in range(250):
        mid = (lo + hi) / 2
        value = npv(mid)
        if abs(value) < 1e-12:
            return mid
        if npv(lo) * value <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def independent_returns(book_path: Path, scenario: str) -> dict[str, float]:
    wb = load_workbook(book_path, data_only=False, read_only=False)
    controls, hist, init = wb["Controls"], wb["Historical & Baseline"], wb["Initiatives"]
    case_row = {controls.cell(r, 3).value: r for r in range(15, 18)}[scenario]
    toggle_rows = {"Pricing and mix": 8, "Commercial efficiency": 9, "AI finance automation": 10}
    revenue = float(hist["F7"].value)
    opening_revenue = revenue
    opening_margin = float(hist["F8"].value) / float(hist["F7"].value)
    opening_debt = float(hist["F8"].value) * float(controls["D22"].value)
    original_debt = opening_debt
    cash = float(controls["D33"].value)
    initial_ebitda = float(hist["F8"].value)
    entry_ev = initial_ebitda * float(controls["D21"].value)
    entry_fees = entry_ev * float(controls["D23"].value)
    sponsor_equity = entry_ev + entry_fees - original_debt
    for year_idx in range(1, 6):
        growth = float(controls.cell(case_row, 3 + year_idx).value)
        revenue = revenue * (1 + growth)
        opening_margin += float(controls.cell(case_row, 9).value)
        core_ebitda = revenue * opening_margin
        initiative_net = 0.0
        for name, rows in INIT_ROWS.items():
            row = rows[scenario]
            enabled = controls.cell(toggle_rows[name], 4).value == "On"
            start_year = int(init.cell(row, 5).value)
            realization = float(init.cell(row, 6).value)
            implementation = float(init.cell(row, 7).value)
            run_cost = float(init.cell(row, 8).value)
            ramp = float(init.cell(row, 10 + year_idx).value)
            active = enabled and year_idx >= start_year
            impl_cost = implementation if enabled and year_idx == start_year else 0.0
            run = run_cost if active else 0.0
            if name == "Pricing and mix":
                benefit = (float(init["D28"].value) - float(init["E28"].value)) * ramp * realization if active else 0.0
            elif name == "Commercial efficiency":
                benefit = float(init["D29"].value) * ramp * realization if active else 0.0
            else:
                capacity = float(init["F30"].value) * ramp if active else 0.0
                gross_value = capacity * float(init["G30"].value) * float(init["H30"].value) / float(controls["D32"].value)
                benefit = gross_value * realization if active else 0.0
            initiative_net += benefit - impl_cost - run
        ebitda = core_ebitda + initiative_net
        da = revenue * float(controls["D27"].value)
        interest = opening_debt * float(controls["D26"].value)
        tax = max(0.0, ebitda - da - interest) * float(controls["D25"].value)
        capex = revenue * float(controls["D28"].value)
        nwc = max(0.0, (revenue - opening_revenue) * float(controls["D29"].value))
        fcf = ebitda - tax - interest - capex - nwc
        mandatory = min(opening_debt, original_debt * float(controls["D30"].value))
        available = cash + fcf - mandatory
        sweep = min(opening_debt - mandatory, max(0.0, available) * float(controls["D31"].value))
        opening_debt = opening_debt - mandatory - sweep
        cash = available - sweep
        opening_revenue = revenue
    exit_multiple = float(controls.cell(case_row, 10).value)
    exit_ev = ebitda * exit_multiple
    exit_fees = exit_ev * float(controls["D24"].value)
    proceeds = exit_ev - exit_fees - opening_debt + cash
    cashflows = [-sponsor_equity] + [float(controls["D33"].value)] * 4 + [proceeds]
    dates = [datetime(2025, 12, 31), datetime(2026, 12, 31), datetime(2027, 12, 31), datetime(2028, 12, 31), datetime(2029, 12, 31), datetime(2030, 12, 31)]
    return {"entry_equity": sponsor_equity, "exit_ebitda": ebitda, "exit_multiple": exit_multiple,
            "exit_ev": exit_ev, "exit_equity": proceeds,
            "exit_debt": opening_debt, "exit_cash": cash, "exit_proceeds": proceeds,
            "moic": proceeds / sponsor_equity, "irr": xirr(cashflows, dates)}


def formula_errors(book_path: Path) -> list[tuple[str, str, object]]:
    wb = load_workbook(book_path, data_only=True, read_only=True)
    issues = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value.startswith(("#REF!", "#DIV/0!", "#VALUE!"))):
                    issues.append((ws.title, cell.coordinate, cell.value))
    return issues


def run_scenario_checks(workbook: Path = WORKBOOK, *, refresh_base: bool = True) -> dict[str, dict[str, float]]:
    outputs: dict[str, dict[str, float]] = {}
    for scenario in SCENARIOS:
        temp = fresh_run_dir(f"scenario_{scenario.lower()}")
        source = build_workbook(temp / f"{scenario}.xlsx", scenario=scenario)
        calculated = recalculate(source, ROOT / "out" / f"calc_{scenario}_{uuid.uuid4().hex[:8]}", ROOT / "out" / f"profile_{scenario}_{uuid.uuid4().hex[:8]}")
        errors = formula_errors(calculated)
        if errors:
            raise AssertionError(f"Formula errors in {scenario}: {errors[:12]}")
        values = load_workbook(calculated, data_only=True, read_only=True)
        if scenario == "Base":
            statuses = [values["Checks & Definitions"].cell(row, 5).value for row in range(8, 20)]
            if statuses != ["PASS"] * len(statuses):
                raise AssertionError(f"Base checks did not all pass: {statuses}")
        computed = independent_returns(source, scenario)
        actual = {"moic": float(values["Returns"]["D25"].value), "irr": float(values["Returns"]["D26"].value)}
        for metric in ("moic", "irr"):
            if not math.isclose(actual[metric], computed[metric], rel_tol=1e-7, abs_tol=1e-7):
                raise AssertionError(f"{scenario} {metric} mismatch: workbook={actual[metric]:.10f}, independent={computed[metric]:.10f}")
        outputs[scenario] = {**computed, **{f"workbook_{k}": v for k, v in actual.items()}}
        if scenario == "Base" and refresh_base:
            shutil.copy2(calculated, workbook)
    if len({round(v["moic"], 8) for v in outputs.values()}) != len(SCENARIOS):
        raise AssertionError("The three scenarios did not produce distinct MOIC results.")
    return outputs


def run_toggle_checks(workbook: Path = WORKBOOK) -> dict[str, float]:
    results = {}
    for idx, (initiative, output_row) in enumerate([
            ("Pricing and mix", 17), ("Commercial efficiency", 23), ("AI finance automation", 32)
        ]):
            temp = fresh_run_dir(f"toggle_{idx}")
            source = temp / f"toggle_{idx}.xlsx"
            source = build_workbook(source, toggles={initiative: "Off"})
            calculated = recalculate(source, ROOT / "out" / f"toggle_calc_{idx}_{uuid.uuid4().hex[:8]}", ROOT / "out" / f"toggle_profile_{idx}_{uuid.uuid4().hex[:8]}")
            values = load_workbook(calculated, data_only=True, read_only=True)
            contribution = sum(float(values["Operating Model"].cell(output_row, col).value or 0) for col in range(4, 9))
            if abs(contribution) > 1e-8:
                raise AssertionError(f"Turning off {initiative} left a contribution of {contribution}.")
            results[initiative] = contribution
    return results


def run_input_variant_check(*, pricing_realization: float | None = None,
                            exit_multiple: float | None = None) -> dict[str, float]:
    """Recalculate one isolated Base-case input change and return its headline outputs."""
    temp = fresh_run_dir("input_variant")
    source = build_workbook(temp / "variant.xlsx", scenario="Base")
    wb = load_workbook(source)
    if pricing_realization is not None:
        wb["Initiatives"]["F7"] = pricing_realization
    if exit_multiple is not None:
        wb["Controls"]["J15"] = exit_multiple
    wb.save(source)
    calculated = recalculate(source, ROOT / "out" / f"variant_calc_{uuid.uuid4().hex[:8]}",
                             ROOT / "out" / f"variant_profile_{uuid.uuid4().hex[:8]}")
    errors = formula_errors(calculated)
    if errors:
        raise AssertionError(f"Formula errors in input variant: {errors[:12]}")
    values = load_workbook(calculated, data_only=True, read_only=True)
    return {"moic": float(values["Returns"]["D25"].value),
            "irr": float(values["Returns"]["D26"].value),
            "exit_ebitda": float(values["Returns"]["D18"].value),
            "exit_equity": float(values["Returns"]["D23"].value)}


def main() -> None:
    parser = argparse.ArgumentParser(description="Recalculate all scenarios and independently verify sponsor returns.")
    parser.add_argument("--no-refresh-base", action="store_true", help="Do not replace the workbook with the recalculated Base copy.")
    args = parser.parse_args()
    scenarios = run_scenario_checks(WORKBOOK, refresh_base=not args.no_refresh_base)
    toggles = run_toggle_checks(WORKBOOK)
    if not args.no_refresh_base:
        base_multiple_down = run_input_variant_check(exit_multiple=9.0)
        base_multiple_up = run_input_variant_check(exit_multiple=11.0)
        pricing_90 = run_input_variant_check(pricing_realization=.90)
    print("Returns ($m except multiple, MOIC and IRR):")
    print("Case      Entry equity   Exit EBITDA   Exit multiple   Exit equity   MOIC    IRR")
    for case, result in scenarios.items():
        print(f"{case:<9} {result['entry_equity']:>12.2f} {result['exit_ebitda']:>13.2f}"
              f" {result['exit_multiple']:>14.1f}x {result['exit_equity']:>13.2f}"
              f" {result['moic']:>6.4f}x {result['irr']:>6.2%} | independent check matched")
    if not args.no_refresh_base:
        values = load_workbook(WORKBOOK, data_only=True, read_only=True)
        sensitivity = values["Sensitivity"]
        print("Sensitivity 1 (MOIC; base F12): "
              f"base {sensitivity['F12'].value:.4f}x; corners "
              f"D10 {sensitivity['D10'].value:.4f}x, H10 {sensitivity['H10'].value:.4f}x, "
              f"D14 {sensitivity['D14'].value:.4f}x, H14 {sensitivity['H14'].value:.4f}x")
        print("Sensitivity 2 (MOIC; base F25): "
              f"base {sensitivity['F25'].value:.4f}x; corners "
              f"D23 {sensitivity['D23'].value:.4f}x, H23 {sensitivity['H23'].value:.4f}x, "
              f"D27 {sensitivity['D27'].value:.4f}x, H27 {sensitivity['H27'].value:.4f}x")
        print(f"Base exit multiple 9.0x: MOIC {base_multiple_down['moic']:.4f}x | IRR {base_multiple_down['irr']:.2%}")
        print(f"Base exit multiple 11.0x: MOIC {base_multiple_up['moic']:.4f}x | IRR {base_multiple_up['irr']:.2%}")
        print(f"Base pricing realization 90%: MOIC {pricing_90['moic']:.4f}x | IRR {pricing_90['irr']:.2%}")
    print("Toggle checks: " + ", ".join(f"{name}=0.0000" for name in toggles))
    print("Formula error scan: no errors in Base, Downside, or Upside.")
    print("Base case workbook checks: all PASS.")


if __name__ == "__main__":
    main()
