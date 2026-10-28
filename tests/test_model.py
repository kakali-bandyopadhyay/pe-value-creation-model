from __future__ import annotations

import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from build.build_model import OUTPUT, build_workbook
from check.verify_returns import run_scenario_checks, run_toggle_checks

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture(scope="session")
def verified_model():
    build_workbook(OUTPUT)
    scenarios = run_scenario_checks(OUTPUT, refresh_base=True)
    toggles = run_toggle_checks(OUTPUT)
    return scenarios, toggles


def test_expected_workbook_tabs_and_live_formula_engine(verified_model):
    wb = load_workbook(OUTPUT, data_only=False)
    assert wb.sheetnames == [
        "Start Here", "Dashboard", "Controls", "Historical & Baseline", "Initiatives",
        "Operating Model", "Cash Flow & Debt", "Returns", "Sensitivity", "VCP Tracker",
        "Checks & Definitions",
    ]
    formulas = [cell for ws in wb.worksheets for row in ws.iter_rows() for cell in row if cell.data_type == "f"]
    assert len(formulas) > 250
    dashboard = wb["Dashboard"]
    assert [dashboard.cell(36, col).value for col in range(7, 12)] == [
        "='Operating Model'!D9", "='Operating Model'!E9", "='Operating Model'!F9",
        "='Operating Model'!G9", "='Operating Model'!H9",
    ]
    # LibreOffice writes a blank calcMode when it saves cached formula results;
    # blank is the spreadsheet default, while the openpyxl source build sets auto.
    assert wb.calculation.calcMode in (None, "auto")
    assert wb.calculation.fullCalcOnLoad is True


def test_scenario_selection_changes_operating_and_returns(verified_model):
    scenarios, _ = verified_model
    assert set(scenarios) == {"Base", "Downside", "Upside"}
    assert len({round(item["moic"], 6) for item in scenarios.values()}) == 3
    assert len({round(item["irr"], 6) for item in scenarios.values()}) == 3


@pytest.mark.parametrize("scenario", ["Base", "Downside", "Upside"])
def test_python_returns_check_matches_recalculated_workbook(verified_model, scenario):
    scenarios, _ = verified_model
    result = scenarios[scenario]
    assert result["moic"] == pytest.approx(result["workbook_moic"], abs=1e-7)
    assert result["irr"] == pytest.approx(result["workbook_irr"], abs=1e-7)


def test_all_base_model_checks_pass(verified_model):
    wb = load_workbook(OUTPUT, data_only=True, read_only=True)
    assert [wb["Checks & Definitions"].cell(row, 5).value for row in range(8, 20)] == ["PASS"] * 12


@pytest.mark.parametrize("initiative", ["Pricing and mix", "Commercial efficiency", "AI finance automation"])
def test_turning_off_a_lever_removes_its_net_contribution(verified_model, initiative):
    _, toggles = verified_model
    assert toggles[initiative] == pytest.approx(0.0, abs=1e-8)


def test_every_scenario_has_a_validated_selector_and_editable_toggles(verified_model):
    wb = load_workbook(OUTPUT, data_only=False)
    ws = wb["Controls"]
    dv_formulas = {dv.formula1 for dv in ws.data_validations.dataValidation}
    assert '"Base,Downside,Upside"' in dv_formulas
    assert '"On,Off"' in dv_formulas
    assert ws["D6"].value == "Base"
    assert all(ws[f"D{row}"].value == "On" for row in (8, 9, 10))


def test_formula_cells_have_no_embedded_business_assumption_literals(verified_model):
    wb = load_workbook(OUTPUT, data_only=False, read_only=True)
    numeric_constants = []
    pattern = re.compile(r"(?<![A-Z0-9_$])(?:\d+(?:\.\d+)?|\.\d+)(?![A-Z0-9_])")
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type != "f":
                    continue
                for token in pattern.findall(cell.value):
                    if float(token) not in (0.0, 1.0):
                        numeric_constants.append((ws.title, cell.coordinate, token, cell.value))
    assert numeric_constants == []


def test_recalculated_scenarios_have_no_formula_errors(verified_model):
    from check.verify_returns import formula_errors
    assert formula_errors(OUTPUT) == []


def test_vcp_tracker_has_quarterly_synthetic_actuals_and_dynamic_statuses(verified_model):
    wb = load_workbook(OUTPUT, data_only=False, read_only=True)
    ws = wb["VCP Tracker"]
    assert ws["C8"].value == "Pricing and mix"
    assert ws["G8"].data_type == "f"
    assert ws["H8"].data_type == "f"
    assert ws["I8"].data_type == "f"
    assert ws["J8"].data_type == "f"
    assert ws["L8"].value == "n.a."
    assert ws["L56"].data_type == "f"  # AI gate is formula-driven in the later-year tracker rows.
    assert ws["L48"].value == '=IF(H48<Controls!$D$39,"Pause",IF(K48<Controls!$D$40,"Fix","Expand"))'
    controls = wb["Controls"]
    assert "does not change the forecast" in controls["D41"].value
    assert "Pause takes precedence" in controls["D41"].value


def test_vcp_tracker_plan_is_non_decreasing_across_each_initiative(verified_model):
    wb = load_workbook(OUTPUT, data_only=True, read_only=True)
    initiatives = wb["Initiatives"]
    tracker = wb["VCP Tracker"]
    for start_row, initiative_row in ((8, 7), (28, 14), (48, 21)):
        annual_ramps = [initiatives.cell(initiative_row, col).value for col in range(11, 16)]
        assert all(current >= previous for previous, current in zip(annual_ramps, annual_ramps[1:]))
        quarterly_plan = [tracker.cell(row, 7).value for row in range(start_row, start_row + 20)]
        assert len(quarterly_plan) == 20
        assert all(current >= previous - 1e-10 for previous, current in zip(quarterly_plan, quarterly_plan[1:]))
        # Year 2 Q1 interpolates 25% of the gap from the prior annual level.
        expected = annual_ramps[0] + (annual_ramps[1] - annual_ramps[0]) * .25
        assert quarterly_plan[4] == pytest.approx(expected)


def test_sensitivity_grids_increase_with_each_positive_driver(verified_model):
    wb = load_workbook(OUTPUT, data_only=True, read_only=True)
    ws = wb["Sensitivity"]
    for min_row, max_row in ((10, 14), (23, 27)):
        values = [[ws.cell(row, col).value for col in range(4, 9)] for row in range(min_row, max_row + 1)]
        assert all(row[i + 1] > row[i] for row in values for i in range(4))
        assert all(values[i + 1][j] > values[i][j] for i in range(4) for j in range(5))
    returns = wb["Returns"]["D25"].value
    assert ws["F12"].value == pytest.approx(returns)
    assert ws["F25"].value == pytest.approx(returns)


def test_content_is_synthetic_and_restricted_terms_are_absent(verified_model):
    fragments = [
        ("Delo", "itte"), ("P", "w", "C"), ("Black", "stone"), ("E", "Q", "T"),
        ("Black", "Rock"), ("K", "K", "R"), ("Car", "lyle"), ("A", "pollo"),
        ("Per", "mira"), ("Morgan", "Stanley"), ("Gold", "man"), ("JP", "Morgan"),
        ("Ci", "ti"), ("Swiss", " Re"),
    ]
    blocked = ["".join(parts) for parts in fragments]
    forbidden_word = "inter" + "view"
    generated_dirs = {".git", ".venv", ".test-deps", ".pip-temp", "__pycache__", ".pytest_cache", "out"}
    for path in ROOT.rglob("*"):
        relative_parts = path.relative_to(ROOT).parts
        if (not path.is_file() or path.suffix.lower() in {".xlsx", ".png", ".pyc"}
                or any(part in generated_dirs or part.startswith("outretry-") for part in relative_parts[:-1])):
            continue
        text = path.read_text(encoding="utf-8", errors="ignore")
        assert forbidden_word.casefold() not in text.casefold(), f"Restricted word in {path.relative_to(ROOT)}"
        assert ("fic" + "tional") not in text.casefold(), f"Use 'synthetic', not the restricted word, in {path.relative_to(ROOT)}"
        for term in blocked:
            assert re.search(rf"\b{re.escape(term)}\b", text, re.IGNORECASE) is None, f"Restricted name in {path.relative_to(ROOT)}"


def test_required_docs_diagram_images_and_artifact_files_exist(verified_model):
    expected = [
        ROOT / "model" / "PE_Value_Creation_Model_synthetic.xlsx",
        ROOT / "docs" / "CASE_STUDY.md",
        ROOT / "docs" / "ARCHITECTURE.md",
        ROOT / "docs" / "dashboard_preview.png",
        ROOT / "docs" / "vcp_tracker_preview.png",
        ROOT / "build" / "build_model.py",
        ROOT / "check" / "verify_returns.py",
    ]
    assert all(path.exists() for path in expected)
    assert "flowchart LR" in (ROOT / "docs" / "ARCHITECTURE.md").read_text(encoding="utf-8")
