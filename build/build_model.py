from __future__ import annotations

"""Rebuild the synthetic PE value-creation model workbook with live formulas."""

from copy import copy
from pathlib import Path
from datetime import date
from openpyxl import Workbook
from openpyxl.chart import LineChart, BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "model" / "PE_Value_Creation_Model_synthetic.xlsx"
NAVY = "17365D"
BLUE = "0000FF"
GREEN = "008000"
BLACK = "202124"
WHITE = "FFFFFF"
YELLOW = "FFF2CC"
PALE_BLUE = "D9EAF7"
LIGHT_BLUE = "EAF2F8"
LIGHT_GREEN = "E2F0D9"
LIGHT_RED = "FCE4D6"
LIGHT_AMBER = "FFF2CC"
GREY = "666666"
LINE = "B7C9D6"
MONEY = '$#,##0.0;($#,##0.0);-'
PCT = '0.0%;(0.0%);-'
MULT = '0.0x'
NUM = '#,##0.0;(#,##0.0);-'
INTEGER = '#,##0;(#,##0);-'

SCENARIOS = ["Base", "Downside", "Upside"]
YEARS = [2026, 2027, 2028, 2029, 2030]
INITIATIVES = ["Pricing and mix", "Commercial efficiency", "AI finance automation"]
OWNER_ROLES = {
    "Pricing and mix": "Commercial workstream owner (synthetic role)",
    "Commercial efficiency": "Operations workstream owner (synthetic role)",
    "AI finance automation": "Finance transformation owner (synthetic role)",
}


def setv(ws, cell, value, *, kind="formula", number_format=None, bold=False,
         fill=None, size=10, align=None, comment=None):
    c = ws[cell]
    c.value = value
    color = BLUE if kind == "input" else (GREEN if isinstance(value, str) and value.startswith("=") and "!" in value else BLACK)
    c.font = Font(name="Arial", size=size, color=color, bold=bold)
    if fill:
        c.fill = PatternFill("solid", fgColor=fill)
    elif kind == "input":
        c.fill = PatternFill("solid", fgColor=YELLOW)
    if number_format:
        c.number_format = number_format
    c.alignment = Alignment(vertical="center", horizontal=align or ("right" if number_format else "left"))
    if comment:
        c.comment = Comment(comment, "Synthetic Model Builder")
    return c


def style_title(ws, title, note, end_col=12):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "C6"
    setv(ws, "C2", title, kind="label", size=14, bold=True)
    setv(ws, "C3", note, kind="label", size=9)
    ws.merge_cells(start_row=3, start_column=3, end_row=3, end_column=end_col)
    ws["C3"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[3].height = 31
    for col in range(3, end_col + 1):
        ws.cell(4, col).border = Border(bottom=Side(style="thin", color=NAVY))
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45


def header(ws, row, start_col, end_col):
    for col in range(start_col, end_col + 1):
        c = ws.cell(row, col)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(bottom=Side(style="thin", color=WHITE))
    ws.row_dimensions[row].height = 30


def section(ws, row, label, end_col=12):
    setv(ws, f"C{row}", label, kind="label", bold=True, fill=PALE_BLUE)
    for col in range(4, end_col + 1):
        ws.cell(row, col).fill = PatternFill("solid", fgColor=PALE_BLUE)
    ws.row_dimensions[row].height = 21


def set_formula_row(ws, row, label, formulas, fmt=None, bold=False):
    setv(ws, f"C{row}", label, kind="label", bold=bold)
    for col, formula in enumerate(formulas, start=4):
        setv(ws, f"{get_column_letter(col)}{row}", formula, number_format=fmt, bold=bold)


def setup_workbook():
    wb = Workbook()
    wb.remove(wb.active)
    names = ["Start Here", "Dashboard", "Controls", "Historical & Baseline", "Initiatives",
             "Operating Model", "Cash Flow & Debt", "Returns", "Sensitivity", "VCP Tracker",
             "Checks & Definitions"]
    for name in names:
        wb.create_sheet(name)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY if ws.title in ("Dashboard", "Returns") else None
        for col in range(1, 18):
            ws.column_dimensions[get_column_letter(col)].width = 14
        ws.column_dimensions["A"].width = 3
        ws.column_dimensions["B"].width = 3
    return wb


def build_start(wb):
    ws = wb["Start Here"]
    style_title(ws, "Synthetic value creation and returns model",
                "All company names, people roles, operating data, targets, and results are synthetic or illustrative. USD millions unless stated. No realized value is claimed.", 11)
    ws.sheet_view.showGridLines = False
    setv(ws, "C6", "NimbleBridge Services (synthetic)", kind="label", bold=True, size=12)
    setv(ws, "C8", "Purpose", kind="label", bold=True)
    setv(ws, "D8", "Explore how a small set of operating initiatives could affect EBITDA, cash generation, debt paydown, and sponsor returns over a five-year ownership period.", kind="label")
    ws["D8"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[8].height = 40
    setv(ws, "C10", "Open Dashboard first. Select a scenario on Controls. Change initiative toggles or assumptions in the yellow cells. Review formula steps on the build tabs and PASS/FAIL checks at the end.", kind="label")
    ws["C10"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("C10:K10")
    ws.row_dimensions[10].height = 40
    setv(ws, "C12", "Workbook navigation", kind="label", bold=True)
    for r, name in enumerate(["Dashboard", "Controls", "Historical & Baseline", "Initiatives", "Operating Model", "Cash Flow & Debt", "Returns", "Sensitivity", "VCP Tracker", "Checks & Definitions"], start=13):
        c = setv(ws, f"C{r}", name, kind="label")
        c.hyperlink = f"#'{name}'!C2"
        c.style = "Hyperlink"
    setv(ws, "F12", "Input and formula legend", kind="label", bold=True)
    for r, label, color, description in [
        (13, "Editable input", YELLOW, "Blue font and pale yellow fill"),
        (14, "Same-sheet formula", WHITE, "Black font"),
        (15, "Linked formula", WHITE, "Green font"),
        (16, "Check", LIGHT_GREEN, "Green PASS or red FAIL"),
    ]:
        setv(ws, f"F{r}", label, kind="label", fill=color)
        setv(ws, f"G{r}", description, kind="label")
    setv(ws, "C25", "Scenario changes operating growth, margin, initiative adoption, timing, realization, implementation costs, and exit multiple. The active forecast uses one selector and one calculation engine.", kind="label")
    ws.merge_cells("C25:K25")
    ws["C25"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[25].height = 40
    for col, width in {"C": 26, "D": 56, "E": 4, "F": 24, "G": 36}.items():
        ws.column_dimensions[col].width = width
    for r in range(13, 23):
        ws.row_dimensions[r].height = 22
    ws.sheet_view.zoomScale = 90


def build_controls(wb):
    ws = wb["Controls"]
    style_title(ws, "Scenario and model controls",
                "All values shown are synthetic assumptions. Change blue-font yellow cells. Currency: USD millions except per-hour cost, hours, and unit conversion.", 10)
    section(ws, 5, "Active case and initiative toggles", 10)
    for row, label, value in [(6, "Selected case", "Base"), (8, "Pricing and mix", "On"), (9, "Commercial efficiency", "On"), (10, "AI finance automation", "On")]:
        setv(ws, f"C{row}", label, kind="label", bold=(row == 6))
        setv(ws, f"D{row}", value, kind="input", bold=(row == 6), comment="Synthetic editable control.")
    dv_case = DataValidation(type="list", formula1='"Base,Downside,Upside"', allow_blank=False)
    dv_onoff = DataValidation(type="list", formula1='"On,Off"', allow_blank=False)
    ws.add_data_validation(dv_case); dv_case.add(ws["D6"])
    ws.add_data_validation(dv_onoff)
    for row in (8, 9, 10): dv_onoff.add(ws[f"D{row}"])
    section(ws, 13, "Scenario assumptions", 10)
    for col, name in enumerate(["Case", "Year 1 growth", "Year 2 growth", "Year 3 growth", "Year 4 growth", "Year 5 growth", "Annual core margin lift", "Exit multiple"], start=3):
        setv(ws, f"{get_column_letter(col)}14", name, kind="label", bold=True, fill=NAVY)
        ws[f"{get_column_letter(col)}14"].font = Font(name="Arial", size=9, bold=True, color=WHITE)
        ws[f"{get_column_letter(col)}14"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    scenario_data = {
        "Base": ([.06, .06, .055, .05, .045], .0025, 10.0),
        "Downside": ([.015, .02, .025, .03, .03], 0.0, 8.5),
        "Upside": ([.08, .075, .07, .06, .05], .004, 11.0),
    }
    for row, case in enumerate(SCENARIOS, start=15):
        setv(ws, f"C{row}", case, kind="input", comment="Synthetic scenario label.")
        for i, rate in enumerate(scenario_data[case][0], start=4):
            setv(ws, f"{get_column_letter(i)}{row}", rate, kind="input", number_format=PCT, comment="Synthetic illustrative operating assumption.")
        setv(ws, f"I{row}", scenario_data[case][1], kind="input", number_format=PCT, comment="Synthetic illustrative annual margin lift.")
        setv(ws, f"J{row}", scenario_data[case][2], kind="input", number_format=MULT, comment="Synthetic illustrative exit valuation assumption.")
    header(ws, 14, 3, 10)
    section(ws, 20, "Transaction, operating, and cash flow assumptions", 10)
    assumptions = [
        (21, "Entry EBITDA multiple", 10.0, MULT),
        (22, "Senior debt / entry EBITDA", 4.0, MULT),
        (23, "Entry transaction fees / entry EV", .02, PCT),
        (24, "Exit fees / exit EV", .01, PCT),
        (25, "Cash tax rate", .25, PCT),
        (26, "Cash interest rate on opening debt", .09, PCT),
        (27, "Depreciation / revenue", .025, PCT),
        (28, "Capital expenditure / revenue", .03, PCT),
        (29, "Working capital investment / revenue change", .08, PCT),
        (30, "Mandatory amortization / original debt", .05, PCT),
        (31, "Cash sweep / positive cash after mandatory amortization", .75, PCT),
        (32, "Dollars per USD million", 1000000, INTEGER),
        (33, "Interim sponsor distribution per forecast year", 0, MONEY),
        (34, "Model check tolerance ($m)", 0.0001, '0.0000'),
        (35, "Required return formula count", 2, INTEGER),
    ]
    for row, label, value, fmt in assumptions:
        setv(ws, f"C{row}", label, kind="label")
        setv(ws, f"D{row}", value, kind="input", number_format=fmt, comment="Synthetic illustrative assumption. Review before changing.")
    section(ws, 36, "Visible value-plan status thresholds", 10)
    for row, label, value, fmt in [
        (37, "Green variance threshold (actual less plan)", -.05, PCT),
        (38, "Amber variance threshold (actual less plan)", -.15, PCT),
        (39, "AI minimum adoption for expand gate", .60, PCT),
        (40, "AI minimum quality for expand gate", .98, PCT),
    ]:
        setv(ws, f"C{row}", label, kind="label")
        setv(ws, f"D{row}", value, kind="input", number_format=fmt, comment="Synthetic illustrative threshold; not an external standard.")
    setv(ws, "C41", "AI gate interpretation", kind="label", bold=True)
    setv(ws, "D41", "Expand / Fix / Pause is a governance signal only; it does not change the forecast. If adoption and quality both fail, Pause takes precedence.", kind="label")
    ws.merge_cells("D41:J41")
    ws["D41"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[41].height = 32
    setv(ws, "C42", "Formula conventions", kind="label", bold=True)
    setv(ws, "D42", "Cash interest uses opening gross debt to avoid circularity. Excess cash is swept at the selected percentage and the remainder stays on balance sheet. No interim sponsor distributions are assumed.", kind="label")
    ws.merge_cells("D42:J42")
    ws["D42"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[42].height = 38
    ws.column_dimensions["C"].width = 42
    for col in range(4, 11): ws.column_dimensions[get_column_letter(col)].width = 19
    ws.freeze_panes = "D15"


def build_history(wb):
    ws = wb["Historical & Baseline"]
    style_title(ws, "Synthetic historical results and entry baseline",
                "Illustrative historical data only. The synthetic company did not report these figures. USD millions.", 8)
    for col, year in enumerate([2023, 2024, 2025], start=4):
        setv(ws, f"{get_column_letter(col)}6", year, kind="label", bold=True, fill=NAVY)
        ws[f"{get_column_letter(col)}6"].font = Font(name="Arial", size=10, bold=True, color=WHITE)
        ws[f"{get_column_letter(col)}6"].alignment = Alignment(horizontal="center")
    header(ws, 6, 4, 6)
    for row, label, values, fmt in [
        (7, "Revenue ($m)", [80.0, 90.0, 100.0], MONEY),
        (8, "EBITDA ($m)", [10.0, 12.2, 14.0], MONEY),
    ]:
        setv(ws, f"C{row}", label, kind="label", bold=(row == 8))
        for col, value in enumerate(values, start=4):
            setv(ws, f"{get_column_letter(col)}{row}", value, kind="input", number_format=fmt,
                 comment="Synthetic illustrative company figure; not an actual company result.")
    setv(ws, "C9", "EBITDA margin", kind="label", bold=True)
    for col in range(4, 7):
        letter = get_column_letter(col)
        setv(ws, f"{letter}9", f"={letter}8/{letter}7", number_format=PCT, bold=True)
    setv(ws, "C12", "Entry baseline is FY2025 revenue and EBITDA. Forecast Year 1 begins in 2026.", kind="label")
    ws.merge_cells("C12:H12")
    ws["C12"].alignment = Alignment(wrap_text=True)
    ws.column_dimensions["C"].width = 44
    for col in range(4, 9): ws.column_dimensions[get_column_letter(col)].width = 18


def scenario_rows():
    return {
        "Pricing and mix": {"Base": 7, "Downside": 8, "Upside": 9},
        "Commercial efficiency": {"Base": 14, "Downside": 15, "Upside": 16},
        "AI finance automation": {"Base": 21, "Downside": 22, "Upside": 23},
    }


def build_initiatives(wb):
    ws = wb["Initiatives"]
    style_title(ws, "Initiative assumptions and execution ramp",
                "All initiative assumptions, roles, KPIs, dependencies, and outcomes are synthetic and illustrative. USD millions unless stated.", 15)
    setv(ws, "C5", "Scenario rows feed the single active forecast. Ramp is the portion of full annual adoption/benefit reached in each ownership year.", kind="label")
    ws.merge_cells("C5:O5")
    ws["C5"].alignment = Alignment(wrap_text=True)
    headers = ["Initiative", "Scenario", "Start year", "Realization", "Implementation cost ($m)", "Annual run cost ($m)", "Confidence", "Dependency", "Year 1 ramp", "Year 2 ramp", "Year 3 ramp", "Year 4 ramp", "Year 5 ramp"]
    for col, label in enumerate(headers, start=3): setv(ws, f"{get_column_letter(col)}6", label, kind="label", bold=True, fill=NAVY)
    header(ws, 6, 3, 15)
    inputs = {
        "Pricing and mix": {
            "Base": (1, .80, .35, .05, .80, "Renewal adoption and churn within guardrail", [.30, .65, .85, .95, 1.00]),
            "Downside": (2, .55, .50, .07, .55, "Renewal adoption and churn within guardrail", [0, .25, .55, .75, .85]),
            "Upside": (1, .90, .32, .04, .90, "Renewal book and sales enablement", [.40, .75, .95, 1.00, 1.00]),
        },
        "Commercial efficiency": {
            "Base": (1, .75, .40, .00, .75, "Spend baseline and accountable cost owners", [.25, .60, .85, 1.00, 1.00]),
            "Downside": (2, .55, .55, .00, .55, "Validated savings and service levels", [0, .25, .50, .70, .80]),
            "Upside": (1, .90, .35, .00, .90, "Procurement and process-owner capacity", [.35, .75, 1.00, 1.00, 1.00]),
        },
        "AI finance automation": {
            "Base": (1, .65, .85, .20, .70, "Data quality, adoption, and process controls", [.25, .55, .75, .85, .90]),
            "Downside": (2, .40, .95, .24, .40, "Quality remediation before workflow expansion", [0, .20, .35, .50, .60]),
            "Upside": (1, .80, .75, .18, .80, "Pilot quality and finance-team adoption", [.35, .70, .90, 1.00, 1.00]),
        },
    }
    for name, rows in scenario_rows().items():
        for case, row in rows.items():
            start, realization, implementation, run_cost, confidence, dependency, ramps = inputs[name][case]
            setv(ws, f"C{row}", name, kind="label")
            setv(ws, f"D{row}", case, kind="input", comment="Synthetic scenario label.")
            setv(ws, f"E{row}", start, kind="input", number_format=INTEGER, comment="Synthetic start year relative to ownership period.")
            setv(ws, f"F{row}", realization, kind="input", number_format=PCT, comment="Synthetic illustrative realization assumption.")
            setv(ws, f"G{row}", implementation, kind="input", number_format=MONEY, comment="Synthetic one-time implementation cost.")
            setv(ws, f"H{row}", run_cost, kind="input", number_format=MONEY, comment="Synthetic annual operating cost from initiative start.")
            setv(ws, f"I{row}", confidence, kind="input", number_format=PCT, comment="Synthetic planning confidence; not a statistical confidence interval.")
            setv(ws, f"J{row}", dependency, kind="input", comment="Synthetic dependency.")
            for col, value in enumerate(ramps, start=11):
                setv(ws, f"{get_column_letter(col)}{row}", value, kind="input", number_format=PCT,
                     comment="Synthetic illustrative ramp by ownership year.")
    section(ws, 26, "Initiative benefit basis", 9)
    for col, label in enumerate(["Initiative", "Full annual gross contribution ($m)", "Churn offset ($m)", "Addressable hours", "Loaded cost ($/hour)", "Avoidable share", "KPI", "Owner role", "Interpretation"], start=3):
        setv(ws, f"{get_column_letter(col)}27", label, kind="label", bold=True, fill=NAVY)
        ws[f"{get_column_letter(col)}27"].font = Font(name="Arial", size=9, bold=True, color=WHITE)
        ws[f"{get_column_letter(col)}27"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    header(ws, 27, 3, 11)
    metadata = [
        ("Pricing and mix", 2.40, .60, 0, 0, 0, "Renewal adoption (%)", OWNER_ROLES["Pricing and mix"], "Gross incremental contribution less churn offset; figures are synthetic."),
        ("Commercial efficiency", 1.80, 0, 0, 0, 0, "Validated savings delivered (%)", OWNER_ROLES["Commercial efficiency"], "Gross annual cost opportunity before realization; figures are synthetic."),
        ("AI finance automation", 0, 0, 18000, 52, .70, "Eligible workflow adoption (%)", OWNER_ROLES["AI finance automation"], "Hours released convert to EBITDA only through avoidable cost share and realization."),
    ]
    for row, data in enumerate(metadata, start=28):
        for col, value in enumerate(data, start=3):
            fmt = MONEY if col in (4, 5) else (INTEGER if col == 6 else (NUM if col == 7 else (PCT if col == 8 else None)))
            setv(ws, f"{get_column_letter(col)}{row}", value, kind="input", number_format=fmt,
                 comment="Synthetic illustrative assumption or role/KPI label.")
    setv(ws, "C33", "Pricing benefit is incremental gross contribution after the displayed churn offset. Commercial savings are treated as EBITDA contributions. AI capacity is converted to avoided contractor or hiring cost only for the displayed avoidable share and realization rate; remaining hours are not valued as EBITDA.", kind="label")
    ws.merge_cells("C33:K34")
    ws["C33"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[33].height = 36
    widths = {"C": 28, "D": 15, "E": 12, "F": 15, "G": 22, "H": 21, "I": 14, "J": 46, "K": 14, "L": 14, "M": 14, "N": 14, "O": 14}
    for col, width in widths.items(): ws.column_dimensions[col].width = width
    ws.freeze_panes = "E7"
    ws.auto_filter.ref = "C6:O23"


def active_lookup(sheet_range, case_range, year=None):
    suffix = f",{year}" if year else ""
    return f"INDEX({sheet_range},MATCH(Controls!$D$6,{case_range},0){suffix})"


def build_operating_model(wb):
    ws = wb["Operating Model"]
    style_title(ws, "Operating model and EBITDA bridge",
                "Selected scenario links to Controls. Annual forecast in USD millions, except capacity hours. All inputs and outputs are synthetic or illustrative.", 10)
    setv(ws, "C5", "Selected case", kind="label", bold=True)
    setv(ws, "D5", "=Controls!D6", number_format=None, bold=True)
    setv(ws, "C6", "Ownership year", kind="label", bold=True)
    for i, col in enumerate(range(4, 9), start=1):
        setv(ws, f"{get_column_letter(col)}6", i, kind="input", number_format=INTEGER, comment="Model period index, not a forecast assumption.")
    header(ws, 6, 3, 8)
    labels = {
        8: "Selected revenue growth", 9: "Core revenue ($m)", 10: "Core EBITDA margin", 11: "Core EBITDA before initiatives ($m)",
        14: "Pricing ramp", 15: "Pricing contribution before costs ($m)", 16: "Pricing implementation and annual run costs ($m)", 17: "Pricing and mix net EBITDA contribution ($m)",
        20: "Commercial efficiency ramp", 21: "Commercial efficiency gross savings ($m)", 22: "Commercial implementation and annual run costs ($m)", 23: "Commercial efficiency net EBITDA contribution ($m)",
        26: "AI workflow adoption ramp", 27: "AI capacity released (hours)", 28: "AI cost value before realization ($m)", 29: "AI avoided-cost savings after realization ($m)", 30: "AI implementation cost ($m)", 31: "AI annual run cost ($m)", 32: "AI finance automation net EBITDA contribution ($m)",
        34: "Net initiative contribution ($m)", 35: "Total EBITDA ($m)", 36: "Total EBITDA margin",
    }
    for row, label in labels.items():
        setv(ws, f"C{row}", label, kind="label", bold=row in (11, 17, 23, 32, 34, 35, 36))
    for row in (7, 12, 13, 18, 19, 24, 25, 33): section(ws, row, {7:"Core forecast",12:"Pricing and mix",13:"Pricing detail",18:"Commercial efficiency",19:"Commercial detail",24:"AI finance automation",25:"AI capacity and benefit",33:"Consolidated EBITDA"}[row], 8)
    scenario_growth = "Controls!$D$15:$H$17"
    case_range = "Controls!$C$15:$C$17"
    for col in range(4, 9):
        L = get_column_letter(col)
        prev = get_column_letter(col - 1)
        yr = f"{L}$6"
        setv(ws, f"{L}8", f"=INDEX({scenario_growth},MATCH(Controls!$D$6,{case_range},0),{yr})", number_format=PCT)
        revenue_formula = f"='Historical & Baseline'!F7*(1+{L}8)" if col == 4 else f"={prev}9*(1+{L}8)"
        setv(ws, f"{L}9", revenue_formula, number_format=MONEY, bold=True)
        margin_lift = f"INDEX(Controls!$I$15:$I$17,MATCH(Controls!$D$6,Controls!$C$15:$C$17,0))"
        margin_start = "'Historical & Baseline'!F9" if col == 4 else f"{prev}10"
        setv(ws, f"{L}10", f"={margin_start}+{margin_lift}", number_format=PCT)
        setv(ws, f"{L}11", f"={L}9*{L}10", number_format=MONEY, bold=True)
        # initiative lookups by scenario within each three-row block
        p = active_lookup("Initiatives!$K$7:$O$9", "Initiatives!$D$7:$D$9", yr)
        p_start = active_lookup("Initiatives!$E$7:$E$9", "Initiatives!$D$7:$D$9")
        p_real = active_lookup("Initiatives!$F$7:$F$9", "Initiatives!$D$7:$D$9")
        p_impl = active_lookup("Initiatives!$G$7:$G$9", "Initiatives!$D$7:$D$9")
        p_run = active_lookup("Initiatives!$H$7:$H$9", "Initiatives!$D$7:$D$9")
        setv(ws, f"{L}14", f"=IF({yr}>={p_start},{p},0)", number_format=PCT)
        setv(ws, f"{L}15", f"=IF(Controls!$D$8=\"On\",('Initiatives'!$D$28-'Initiatives'!$E$28)*{L}14*{p_real},0)", number_format=MONEY)
        setv(ws, f"{L}16", f"=IF(AND(Controls!$D$8=\"On\",{yr}={p_start}),{p_impl},0)+IF(AND(Controls!$D$8=\"On\",{yr}>={p_start}),{p_run},0)", number_format=MONEY)
        setv(ws, f"{L}17", f"={L}15-{L}16", number_format=MONEY, bold=True)
        c = active_lookup("Initiatives!$K$14:$O$16", "Initiatives!$D$14:$D$16", yr)
        c_start = active_lookup("Initiatives!$E$14:$E$16", "Initiatives!$D$14:$D$16")
        c_real = active_lookup("Initiatives!$F$14:$F$16", "Initiatives!$D$14:$D$16")
        c_impl = active_lookup("Initiatives!$G$14:$G$16", "Initiatives!$D$14:$D$16")
        c_run = active_lookup("Initiatives!$H$14:$H$16", "Initiatives!$D$14:$D$16")
        setv(ws, f"{L}20", f"=IF({yr}>={c_start},{c},0)", number_format=PCT)
        setv(ws, f"{L}21", f"=IF(Controls!$D$9=\"On\",'Initiatives'!$D$29*{L}20*{c_real},0)", number_format=MONEY)
        setv(ws, f"{L}22", f"=IF(AND(Controls!$D$9=\"On\",{yr}={c_start}),{c_impl},0)+IF(AND(Controls!$D$9=\"On\",{yr}>={c_start}),{c_run},0)", number_format=MONEY)
        setv(ws, f"{L}23", f"={L}21-{L}22", number_format=MONEY, bold=True)
        a = active_lookup("Initiatives!$K$21:$O$23", "Initiatives!$D$21:$D$23", yr)
        a_start = active_lookup("Initiatives!$E$21:$E$23", "Initiatives!$D$21:$D$23")
        a_real = active_lookup("Initiatives!$F$21:$F$23", "Initiatives!$D$21:$D$23")
        a_impl = active_lookup("Initiatives!$G$21:$G$23", "Initiatives!$D$21:$D$23")
        a_run = active_lookup("Initiatives!$H$21:$H$23", "Initiatives!$D$21:$D$23")
        setv(ws, f"{L}26", f"=IF({yr}>={a_start},{a},0)", number_format=PCT)
        setv(ws, f"{L}27", f"=IF(Controls!$D$10=\"On\",'Initiatives'!$F$30*{L}26,0)", number_format=INTEGER)
        setv(ws, f"{L}28", f"={L}27*'Initiatives'!$G$30*'Initiatives'!$H$30/Controls!$D$32", number_format=MONEY)
        setv(ws, f"{L}29", f"=IF(Controls!$D$10=\"On\",{L}28*{a_real},0)", number_format=MONEY)
        setv(ws, f"{L}30", f"=IF(AND(Controls!$D$10=\"On\",{yr}={a_start}),{a_impl},0)", number_format=MONEY)
        setv(ws, f"{L}31", f"=IF(AND(Controls!$D$10=\"On\",{yr}>={a_start}),{a_run},0)", number_format=MONEY)
        setv(ws, f"{L}32", f"={L}29-{L}30-{L}31", number_format=MONEY, bold=True)
        setv(ws, f"{L}34", f"=SUM({L}17,{L}23,{L}32)", number_format=MONEY, bold=True)
        setv(ws, f"{L}35", f"={L}11+{L}34", number_format=MONEY, bold=True)
        setv(ws, f"{L}36", f"={L}35/{L}9", number_format=PCT, bold=True)
    for row in (11, 17, 23, 32, 34, 35):
        for col in range(3, 9): ws.cell(row, col).border = Border(top=Side(style="thin", color=LINE))
    for col, width in {"C": 61, "D": 15, "E": 15, "F": 15, "G": 15, "H": 15}.items(): ws.column_dimensions[col].width = width
    ws.freeze_panes = "D7"


def build_cashflow(wb):
    ws = wb["Cash Flow & Debt"]
    style_title(ws, "Cash flow, debt paydown, and retained cash",
                "Selected scenario, annual periods, USD millions. Interest is calculated on opening gross debt to avoid circular formulas. All figures are synthetic.", 9)
    setv(ws, "C5", "Selected case", kind="label", bold=True); setv(ws, "D5", "=Controls!D6", bold=True)
    for i, col in enumerate(range(4, 9), start=1): setv(ws, f"{get_column_letter(col)}6", i, kind="input", number_format=INTEGER, comment="Model period index.")
    header(ws, 6, 3, 8)
    rows = {
        8:"Total EBITDA ($m)",9:"Depreciation ($m)",10:"EBIT ($m)",11:"Opening gross debt ($m)",12:"Cash interest ($m)",13:"Taxable earnings after interest ($m)",14:"Cash tax ($m)",15:"Capital expenditure ($m)",16:"Working capital investment ($m)",17:"Cash flow before debt service ($m)",18:"Opening cash ($m)",19:"Mandatory amortization ($m)",20:"Cash available after mandatory amortization ($m)",21:"Cash sweep ($m)",22:"Ending gross debt ($m)",23:"Ending cash ($m)",24:"Net debt at year-end ($m)"
    }
    for row,label in rows.items(): setv(ws,f"C{row}",label,kind="label",bold=row in (17,22,23,24))
    for col in range(4,9):
        L=get_column_letter(col); prev=get_column_letter(col-1)
        setv(ws,f"{L}8",f"='Operating Model'!{L}35",number_format=MONEY)
        setv(ws,f"{L}9",f"='Operating Model'!{L}9*Controls!$D$27",number_format=MONEY)
        setv(ws,f"{L}10",f"={L}8-{L}9",number_format=MONEY)
        setv(ws,f"{L}11",f"='Returns'!$D$15" if col==4 else f"={prev}22",number_format=MONEY)
        setv(ws,f"{L}12",f"={L}11*Controls!$D$26",number_format=MONEY)
        setv(ws,f"{L}13",f"=MAX(0,{L}10-{L}12)",number_format=MONEY)
        setv(ws,f"{L}14",f"={L}13*Controls!$D$25",number_format=MONEY)
        setv(ws,f"{L}15",f"='Operating Model'!{L}9*Controls!$D$28",number_format=MONEY)
        prior_rev="'Historical & Baseline'!$F$7" if col==4 else f"'Operating Model'!{prev}9"
        setv(ws,f"{L}16",f"=MAX(0,('Operating Model'!{L}9-{prior_rev})*Controls!$D$29)",number_format=MONEY)
        setv(ws,f"{L}17",f"={L}8-{L}14-{L}12-{L}15-{L}16",number_format=MONEY,bold=True)
        setv(ws,f"{L}18",f"=Controls!$D$33" if col==4 else f"={prev}23",number_format=MONEY)
        setv(ws,f"{L}19",f"=MIN({L}11,'Returns'!$D$15*Controls!$D$30)",number_format=MONEY)
        setv(ws,f"{L}20",f"={L}18+{L}17-{L}19",number_format=MONEY)
        setv(ws,f"{L}21",f"=MIN({L}11-{L}19,MAX(0,{L}20)*Controls!$D$31)",number_format=MONEY)
        setv(ws,f"{L}22",f"={L}11-{L}19-{L}21",number_format=MONEY,bold=True)
        setv(ws,f"{L}23",f"={L}20-{L}21",number_format=MONEY,bold=True)
        setv(ws,f"{L}24",f"={L}22-{L}23",number_format=MONEY,bold=True)
    section(ws,27,"Debt and cash definitions",8)
    setv(ws,"C28","Ending gross debt equals opening debt less mandatory amortization and cash sweep. Ending cash equals opening cash plus cash flow before debt service, less mandatory amortization and cash sweep. Net debt equals gross debt less cash.",kind="label")
    ws.merge_cells("C28:H29");ws["C28"].alignment=Alignment(wrap_text=True,vertical="top")
    for col,width in {"C":72,"D":16,"E":16,"F":16,"G":16,"H":16}.items():ws.column_dimensions[col].width=width
    ws.freeze_panes="D7"


def build_returns(wb):
    ws=wb["Returns"]
    style_title(ws,"Entry-to-exit sponsor returns","One selected scenario. Annual dated cash flows in USD millions. Returns are formula-driven and illustrative.",9)
    setv(ws,"C5","Selected case",kind="label",bold=True);setv(ws,"D5","=Controls!D6",bold=True)
    setv(ws,"C7","Cash flow date",kind="label",bold=True);setv(ws,"C8","Sponsor cash flow ($m)",kind="label",bold=True)
    for col,dt in enumerate([date(2025,12,31),date(2026,12,31),date(2027,12,31),date(2028,12,31),date(2029,12,31),date(2030,12,31)],start=4):
        L=get_column_letter(col);setv(ws,f"{L}7",dt,kind="input",number_format="dd-mmm-yyyy",comment="Synthetic illustrative model date.")
    header(ws,7,3,9)
    setv(ws,"D8","=-D16",number_format=MONEY,bold=True)
    for col in range(5,9):setv(ws,f"{get_column_letter(col)}8","=Controls!$D$33",number_format=MONEY)
    setv(ws,"I8","=D23",number_format=MONEY,bold=True)
    items=[
        (11,"Entry EBITDA ($m)","='Historical & Baseline'!F8",MONEY),
        (12,"Entry EBITDA multiple","=Controls!$D$21",MULT),
        (13,"Entry enterprise value ($m)","=D11*D12",MONEY),
        (14,"Entry transaction fees ($m)","=D13*Controls!$D$23",MONEY),
        (15,"Senior debt at entry ($m)","=D11*Controls!$D$22",MONEY),
        (16,"Sponsor equity at entry ($m)","=D13+D14-D15",MONEY),
        (18,"Exit EBITDA ($m)","='Operating Model'!H35",MONEY),
        (19,"Exit EBITDA multiple","=INDEX(Controls!$J$15:$J$17,MATCH(Controls!$D$6,Controls!$C$15:$C$17,0))",MULT),
        (20,"Exit enterprise value ($m)","=D18*D19",MONEY),
        (21,"Exit fees ($m)","=D20*Controls!$D$24",MONEY),
        (22,"Exit net debt: gross debt less cash ($m)","='Cash Flow & Debt'!H24",MONEY),
        (23,"Exit equity proceeds ($m)","=D20-D21-D22",MONEY),
        (25,"Sponsor MOIC","=D23/D16",MULT),
        (26,"Sponsor IRR from dated cash flows","=XIRR(D8:I8,D7:I7)",PCT),
    ]
    for row,label,formula,fmt in items:
        setv(ws,f"C{row}",label,kind="label",bold=row in (16,23,25,26))
        setv(ws,f"D{row}",formula,number_format=fmt,bold=row in (16,23,25,26))
    section(ws,29,"Returns calculation notes",8)
    notes=["Entry enterprise value = FY2025 EBITDA × entry EBITDA multiple.","Sponsor equity funds enterprise value plus entry fees less senior debt at entry.","Exit equity proceeds = exit enterprise value less exit fees and exit net debt (gross debt less retained cash).","MOIC = exit equity proceeds ÷ sponsor equity at entry. Interim sponsor distributions are shown in the dated cash-flow row.","IRR uses XIRR over the dated sponsor cash flows. With no interim distributions, the only positive flow is at exit."]
    for i,note in enumerate(notes,start=30):setv(ws,f"C{i}",note,kind="label")
    ws.merge_cells("C30:H34");ws["C30"].alignment=Alignment(wrap_text=True,vertical="top")
    ws.column_dimensions["C"].width=48
    for col in range(4,10):ws.column_dimensions[get_column_letter(col)].width=17
    ws.freeze_panes="D8"


def build_sensitivity(wb):
    ws=wb["Sensitivity"]
    style_title(ws,"Returns sensitivity","Formula grids. The first grid holds selected-case net debt constant. The second recomputes operating cash, debt paydown, and retained cash for each AI realization input. USD millions.",11)
    setv(ws,"C5","Selected case",kind="label",bold=True);setv(ws,"D5","=Controls!D6",bold=True)
    section(ws,7,"Exit multiple × exit EBITDA, MOIC (selected-case net debt held fixed)",10)
    setv(ws,"C8","Direction: MOIC rises left to right with exit multiple and top to bottom with exit EBITDA.",kind="label")
    ws.merge_cells("C8:H8")
    setv(ws,"C9","Exit EBITDA ($m)",kind="label",bold=True,fill=NAVY)
    setv(ws,"J9","Exit EBITDA factor",kind="label",bold=True,fill=NAVY)
    ws["J9"].font=Font(name="Arial",size=9,bold=True,color=WHITE)
    ws["J9"].alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
    for col,m in enumerate([8.0,9.0,10.0,11.0,12.0],start=4):setv(ws,f"{get_column_letter(col)}9",m,kind="input",number_format=MULT,comment="Synthetic sensitivity point.")
    header(ws,9,3,8)
    for row,factor in enumerate([.85,.925,1.0,1.075,1.15],start=10):
        setv(ws,f"J{row}",factor,kind="input",number_format=PCT,comment="Synthetic illustrative exit EBITDA sensitivity factor.")
        setv(ws,f"C{row}",f"='Returns'!$D$18*$J{row}",number_format=MONEY)
        for col in range(4,9):
            L=get_column_letter(col)
            setv(ws,f"{L}{row}",f"=(($C{row}*{L}$9)*(1-Controls!$D$24)-'Returns'!$D$22)/'Returns'!$D$16",number_format=MULT)
    setv(ws,"C16","At each grid point, exit enterprise value equals exit EBITDA × exit multiple. Exit fees and selected-case net debt are then deducted before dividing by sponsor equity.",kind="label")
    ws.merge_cells("C16:H17");ws["C16"].alignment=Alignment(wrap_text=True,vertical="top")
    section(ws,20,"Exit multiple × AI cost-realization assumption, MOIC",10)
    setv(ws,"C21","Direction: MOIC rises with exit multiple and AI cost realization.",kind="label")
    ws.merge_cells("C21:H21")
    setv(ws,"C22","Exit multiple",kind="label",bold=True,fill=NAVY)
    realization_points=[.30,.50,.65,.80,1.00]
    for col,rate in enumerate(realization_points,start=4):setv(ws,f"{get_column_letter(col)}22",rate,kind="input",number_format=PCT,comment="Synthetic sensitivity point for AI avoided-cost realization.")
    header(ws,22,3,8)
    for row,multiple in enumerate([8.0,9.0,10.0,11.0,12.0],start=23):
        setv(ws,f"C{row}",multiple,kind="input",number_format=MULT,comment="Synthetic sensitivity point.")
    # Recalculation engines: each realization point recalculates all five annual cash and debt flows.
    section(ws,31,"AI realization sensitivity calculation detail",11)
    setv(ws,"C33","AI cost realization",kind="label",bold=True)
    setv(ws,"C34","Synthetic axis values",kind="label")
    engine_blocks=[]
    for k,rate in enumerate(realization_points):
        base_row=36+k*16
        setv(ws,f"C{base_row}",f"AI realization case {k+1}",kind="label",bold=True)
        setv(ws,f"D{base_row}",f"=D{22+k*0}" if False else rate,kind="input",number_format=PCT,comment="Synthetic sensitivity driver; varies the AI initiative realization only.")
        # yearly schedule header
        for j,col in enumerate(range(5,10),start=1):setv(ws,f"{get_column_letter(col)}{base_row}",j,kind="input",number_format=INTEGER)
        labels={1:"Adjusted total EBITDA ($m)",2:"Depreciation ($m)",3:"Interest ($m)",4:"Cash taxes ($m)",5:"Capital expenditure ($m)",6:"Working capital investment ($m)",7:"Cash flow before debt service ($m)",8:"Opening debt ($m)",9:"Opening cash ($m)",10:"Mandatory amortization ($m)",11:"Cash available after mandatory amortization ($m)",12:"Cash sweep ($m)",13:"Ending gross debt ($m)"}
        for offset,label in labels.items():setv(ws,f"C{base_row+offset}",label,kind="label")
        # End cash and net debt are placed in the adjacent year-end rows on separate labeled lines.
        for j,col in enumerate(range(5,10),start=1):
            L=get_column_letter(col);prev=get_column_letter(col-1);op_col=get_column_letter(col-1);op_prev=get_column_letter(col-2)
            setv(ws,f"{L}{base_row+1}",f"='Operating Model'!{op_col}35-'Operating Model'!{op_col}29+'Operating Model'!{op_col}28*$D{base_row}",number_format=MONEY)
            setv(ws,f"{L}{base_row+2}",f"='Operating Model'!{op_col}9*Controls!$D$27",number_format=MONEY)
            debt_begin=f"'Returns'!$D$15" if j==1 else f"{prev}{base_row+13}"
            cash_begin="Controls!$D$33" if j==1 else f"{prev}{base_row+14}"
            setv(ws,f"{L}{base_row+3}",f"={debt_begin}*Controls!$D$26",number_format=MONEY)
            setv(ws,f"{L}{base_row+4}",f"=MAX(0,{L}{base_row+1}-{L}{base_row+2}-{L}{base_row+3})*Controls!$D$25",number_format=MONEY)
            setv(ws,f"{L}{base_row+5}",f"='Operating Model'!{op_col}9*Controls!$D$28",number_format=MONEY)
            prev_rev="'Historical & Baseline'!$F$7" if j==1 else f"'Operating Model'!{op_prev}9"
            setv(ws,f"{L}{base_row+6}",f"=MAX(0,('Operating Model'!{op_col}9-{prev_rev})*Controls!$D$29)",number_format=MONEY)
            setv(ws,f"{L}{base_row+7}",f"={L}{base_row+1}-{L}{base_row+4}-{L}{base_row+3}-{L}{base_row+5}-{L}{base_row+6}",number_format=MONEY)
            setv(ws,f"{L}{base_row+8}",f"={debt_begin}",number_format=MONEY)
            setv(ws,f"{L}{base_row+9}",f"={cash_begin}",number_format=MONEY)
            setv(ws,f"{L}{base_row+10}",f"=MIN({L}{base_row+8},'Returns'!$D$15*Controls!$D$30)",number_format=MONEY)
            setv(ws,f"{L}{base_row+11}",f"={L}{base_row+9}+{L}{base_row+7}-{L}{base_row+10}",number_format=MONEY)
            setv(ws,f"{L}{base_row+12}",f"=MIN({L}{base_row+8}-{L}{base_row+10},MAX(0,{L}{base_row+11})*Controls!$D$31)",number_format=MONEY)
            setv(ws,f"{L}{base_row+13}",f"={L}{base_row+8}-{L}{base_row+10}-{L}{base_row+12}",number_format=MONEY)
            setv(ws,f"{L}{base_row+14}",f"={L}{base_row+11}-{L}{base_row+12}",number_format=MONEY)
            setv(ws,f"{L}{base_row+15}",f"={L}{base_row+13}-{L}{base_row+14}",number_format=MONEY)
        engine_blocks.append({"base":base_row,"ai_gross":base_row+1,"end_debt":base_row+13,"end_cash":base_row+14,"net_debt":base_row+15})
    # Sensitivity grid rows link to terminal adjusted EBITDA, debt, and cash from each matching engine.
    for row in range(23,28):
        for col in range(4,9):
            L=get_column_letter(col);engine=engine_blocks[col-4];base=engine["base"]
            # Exit EBITDA is selected-case base EBITDA less original AI savings plus the alternative savings.
            setv(ws,f"{L}{row}",f"=(('Operating Model'!$H$35-'Operating Model'!$H$29+'Operating Model'!$H$28*$D${base})*$C{row}*(1-Controls!$D$24)-$I${base+13}+$I${base+14})/'Returns'!$D$16",number_format=MULT)
    setv(ws,"C29","This grid changes only AI cost realization. It recalculates annual taxes, cash flow, mandatory amortization, cash sweep, ending debt, and retained cash at each realization point. Other initiative assumptions stay at the selected case. The 65% column and 10.0x row are the selected Base assumptions.",kind="label")
    ws.merge_cells("C29:H30");ws["C29"].alignment=Alignment(wrap_text=True,vertical="top")
    for col,width in {"C":72,"D":17,"E":16,"F":16,"G":16,"H":16,"I":16}.items():ws.column_dimensions[col].width=width
    ws.freeze_panes="D10"


def build_vcp(wb):
    ws=wb["VCP Tracker"]
    style_title(ws,"Quarterly value creation plan tracker","Synthetic plan and synthetic actuals illustrate ownership, execution measures, variance, RAG status, and next actions. USD unless KPI states percent.",17)
    setv(ws,"C5","RAG thresholds: Green at or above the green variance limit; Amber between the green and amber limits; Red below the amber limit.",kind="label")
    ws.merge_cells("C5:Q5");ws["C5"].alignment=Alignment(wrap_text=True)
    setv(ws,"C6","AI Expand / Fix / Pause is a governance signal only and does not change the forecast. If adoption and quality both fail, Pause takes precedence.",kind="label")
    ws.merge_cells("C6:Q6");ws["C6"].alignment=Alignment(wrap_text=True)
    ws.row_dimensions[5].height=24;ws.row_dimensions[6].height=24
    for col,label in enumerate(["Initiative","Quarter","Owner role","KPI","Plan","Synthetic actual","Variance (actual less plan)","RAG","Quality KPI","AI gate","Next action","Quarter phasing","Synthetic actual input","Year index","RAG threshold note"],start=3):
        setv(ws,f"{get_column_letter(col)}7",label,kind="label",bold=True,fill=NAVY)
    header(ws,7,3,17)
    data=[]
    for initiative, base_row, toggle, kpi in [
        ("Pricing and mix",7,8,"Renewal adoption (% of eligible renewals)"),
        ("Commercial efficiency",14,9,"Savings delivered (% of run-rate target)"),
        ("AI finance automation",21,10,"Workflow adoption (% of eligible workflows)"),
    ]:
        for year in range(1,6):
            for q,phase in enumerate([.25,.50,.75,1.0],start=1):
                # reproducible illustrative synthetic tracking inputs, not actual company performance
                attainment = max(.78, min(1.03, .91 + .025*((year+q) % 5) - .018*(q == 1)))
                quality = (0.975 + .004*((year+q) % 5)) if initiative=="AI finance automation" else None
                data.append((initiative,year,q,phase,attainment,quality,base_row,kpi))
    for row,item in enumerate(data,start=8):
        initiative,year,q,phase,attainment,quality,base_row,kpi=item
        setv(ws,f"C{row}",initiative,kind="label")
        setv(ws,f"D{row}",f"{YEARS[year-1]} Q{q}",kind="input",comment="Synthetic period label.")
        setv(ws,f"E{row}",OWNER_ROLES[initiative],kind="input",comment="Synthetic accountable owner role.")
        setv(ws,f"F{row}",kpi,kind="input",comment="Synthetic KPI definition.")
        ramp_col=get_column_letter(10+year)
        setv(ws,f"N{row}",phase,kind="input",number_format=PCT,comment="Synthetic quarterly phasing of annual adoption/ramp.")
        setv(ws,f"O{row}",attainment,kind="input",number_format=PCT,comment="Synthetic actual attainment factor against plan.")
        setv(ws,f"P{row}",year,kind="input",number_format=INTEGER,comment="Model ownership year index.")
        if year == 1:
            plan_formula=f"=Initiatives!{ramp_col}{base_row}*N{row}"
        else:
            prior_ramp_col=get_column_letter(10+year-1)
            plan_formula=f"=Initiatives!{prior_ramp_col}{base_row}+(Initiatives!{ramp_col}{base_row}-Initiatives!{prior_ramp_col}{base_row})*N{row}"
        setv(ws,f"G{row}",plan_formula,number_format=PCT)
        setv(ws,f"H{row}",f"=G{row}*O{row}",number_format=PCT)
        setv(ws,f"I{row}",f"=H{row}-G{row}",number_format=PCT)
        setv(ws,f"J{row}",f'=IF(I{row}>=Controls!$D$37,"Green",IF(I{row}>=Controls!$D$38,"Amber","Red"))',kind="formula")
        if quality is not None:
            setv(ws,f"K{row}",quality,kind="input",number_format=PCT,comment="Synthetic AI output-quality measure; illustrative only.")
            setv(ws,f"L{row}",f'=IF(H{row}<Controls!$D$39,"Pause",IF(K{row}<Controls!$D$40,"Fix","Expand"))')
        else:
            setv(ws,f"K{row}","n.a.",kind="label")
            setv(ws,f"L{row}","n.a.",kind="label")
        action = "Review adoption barrier and next renewal cohort" if initiative=="Pricing and mix" else ("Confirm validated savings and service guardrails" if initiative=="Commercial efficiency" else "Expand only after adoption and quality gates; otherwise fix or pause")
        setv(ws,f"M{row}",action,kind="label")
        setv(ws,f"Q{row}","Synthetic plan and sample tracking; not measured company results",kind="label")
    ws.conditional_formatting.add(f"J8:J{7+len(data)}",FormulaRule(formula=['J8="Green"'],fill=PatternFill("solid",fgColor=LIGHT_GREEN)))
    ws.conditional_formatting.add(f"J8:J{7+len(data)}",FormulaRule(formula=['J8="Amber"'],fill=PatternFill("solid",fgColor=LIGHT_AMBER)))
    ws.conditional_formatting.add(f"J8:J{7+len(data)}",FormulaRule(formula=['J8="Red"'],fill=PatternFill("solid",fgColor=LIGHT_RED)))
    ws.conditional_formatting.add(f"L8:L{7+len(data)}",FormulaRule(formula=['L8="Expand"'],fill=PatternFill("solid",fgColor=LIGHT_GREEN)))
    ws.conditional_formatting.add(f"L8:L{7+len(data)}",FormulaRule(formula=['L8="Fix"'],fill=PatternFill("solid",fgColor=LIGHT_AMBER)))
    ws.conditional_formatting.add(f"L8:L{7+len(data)}",FormulaRule(formula=['L8="Pause"'],fill=PatternFill("solid",fgColor=LIGHT_RED)))
    ws.auto_filter.ref=f"C7:Q{7+len(data)}"
    ws.freeze_panes="G8"
    widths={"C":27,"D":14,"E":42,"F":36,"G":12,"H":14,"I":18,"J":12,"K":15,"L":12,"M":56,"N":15,"O":18,"P":12,"Q":48}
    for col,width in widths.items():ws.column_dimensions[col].width=width
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0


def build_dashboard(wb):
    ws=wb["Dashboard"]
    style_title(ws,"Value creation plan and sponsor returns","NimbleBridge Services (synthetic). All values and plan tracking are illustrative. Selected scenario results update from workbook controls.",13)
    setv(ws,"C5","Selected case",kind="label",bold=True);setv(ws,"D5","=Controls!D6",bold=True)
    section(ws,7,"Return snapshot",9)
    for row,label,formula,fmt in [(8,"Entry EBITDA ($m)","=Returns!D11",MONEY),(9,"Entry enterprise value ($m)","=Returns!D13",MONEY),(10,"Sponsor equity at entry ($m)","=Returns!D16",MONEY),(11,"Exit EBITDA ($m)","=Returns!D18",MONEY),(12,"Exit enterprise value ($m)","=Returns!D20",MONEY),(13,"Exit net debt ($m)","='Cash Flow & Debt'!H24",MONEY),(14,"Sponsor MOIC","=Returns!D25",MULT),(15,"Sponsor IRR","=Returns!D26",PCT)]:
        setv(ws,f"C{row}",label,kind="label",bold=row in (14,15));setv(ws,f"D{row}",formula,number_format=fmt,bold=row in (14,15))
    section(ws,18,"EBITDA bridge, FY2025 baseline to FY2030 exit",10)
    bridge=[(19,"FY2025 baseline EBITDA ($m)","='Historical & Baseline'!F8"),(20,"Core revenue and margin change ($m)","='Operating Model'!H11-'Historical & Baseline'!F8"),(21,"Pricing and mix contribution ($m)","='Operating Model'!H17"),(22,"Commercial efficiency contribution ($m)","='Operating Model'!H23"),(23,"AI finance automation contribution ($m)","='Operating Model'!H32"),(24,"FY2030 exit EBITDA ($m)","='Operating Model'!H35")]
    for row,label,formula in bridge:setv(ws,f"C{row}",label,kind="label",bold=row in (19,24));setv(ws,f"D{row}",formula,number_format=MONEY,bold=row in (19,24))
    section(ws,27,"Enterprise value bridge at entry multiple",10)
    for row,label,formula in [(28,"Entry enterprise value ($m)","=Returns!D13"),(29,"EBITDA change at entry multiple ($m)","=(Returns!D18-Returns!D11)*Returns!D12"),(30,"Exit multiple change ($m)","=Returns!D18*(Returns!D19-Returns!D12)"),(31,"Exit enterprise value ($m)","=SUM(D28:D30)")]:
        setv(ws,f"C{row}",label,kind="label",bold=row in (28,31));setv(ws,f"D{row}",formula,number_format=MONEY,bold=row in (28,31))
    section(ws,34,"Operating trend",10)
    for col,year in enumerate([2023,2024,2025,2026,2027,2028,2029,2030],start=4):setv(ws,f"{get_column_letter(col)}35",year,kind="label",bold=True,fill=NAVY);ws[f"{get_column_letter(col)}35"].font=Font(name="Arial",size=9,bold=True,color=WHITE)
    header(ws,35,3,11)
    setv(ws,"C36","Revenue ($m)",kind="label");setv(ws,"C37","EBITDA ($m)",kind="label")
    for col in range(4,12):
        L=get_column_letter(col)
        if col<=6:
            hist_col=get_column_letter(col)
            setv(ws,f"{L}36",f"='Historical & Baseline'!{hist_col}7",number_format=MONEY)
            setv(ws,f"{L}37",f"='Historical & Baseline'!{hist_col}8",number_format=MONEY)
        else:
            op_col=get_column_letter(col-3)
            setv(ws,f"{L}36",f"='Operating Model'!{op_col}9",number_format=MONEY)
            setv(ws,f"{L}37",f"='Operating Model'!{op_col}35",number_format=MONEY)
    chart=LineChart();chart.title="Revenue and EBITDA";chart.y_axis.title="USD millions";chart.x_axis.title="Year";chart.height=7;chart.width=14
    chart.add_data(Reference(ws,min_col=3,max_col=11,min_row=36,max_row=37),from_rows=True,titles_from_data=True)
    chart.set_categories(Reference(ws,min_col=4,max_col=11,min_row=35));chart.legend.position="b";ws.add_chart(chart,"F7")
    # Summary table links to VCP record-level formulas and preserves denominator/count transparency.
    section(ws,40,"Value plan tracker summary",11)
    for row,label,formula in [(41,"Quarterly initiative records","=COUNTA('VCP Tracker'!C8:C67)"),(42,"Green RAG records","=COUNTIF('VCP Tracker'!J8:J67,\"Green\")"),(43,"Amber RAG records","=COUNTIF('VCP Tracker'!J8:J67,\"Amber\")"),(44,"Red RAG records","=COUNTIF('VCP Tracker'!J8:J67,\"Red\")"),(45,"AI expand-gate records","=COUNTIF('VCP Tracker'!L8:L67,\"Expand\")")]:setv(ws,f"C{row}",label,kind="label");setv(ws,f"D{row}",formula,number_format=INTEGER)
    section(ws,48,"Selected-case sensitivities",10)
    setv(ws,"C49","MOIC by exit EBITDA and multiple",kind="label",bold=True)
    for row in range(10,15):
        for col in range(4,9):setv(ws,f"{get_column_letter(col)}{row+40}",f"=Sensitivity!{get_column_letter(col)}{row}",number_format=MULT)
    setv(ws,"C50","MOIC by AI realization and exit multiple",kind="label",bold=True)
    for row in range(23,28):
        for col in range(4,9):setv(ws,f"{get_column_letter(col)}{row+28}",f"=Sensitivity!{get_column_letter(col)}{row}",number_format=MULT)
    setv(ws,"C57","Risks and interpretation",kind="label",bold=True,fill=PALE_BLUE)
    risk_text="Synthetic inputs are not diligence evidence. Adoption, quality, churn, savings realization, margin, cash tax, working capital, financing, and exit valuation need independent validation. Exit EBITDA changes affect returns; the value bridge separates EBITDA improvement from multiple change. No value is represented as realized."
    setv(ws,"C58",risk_text,kind="label");ws.merge_cells("C58:K60");ws["C58"].alignment=Alignment(wrap_text=True,vertical="top")
    ws.column_dimensions["C"].width=49
    for col in range(4,12):ws.column_dimensions[get_column_letter(col)].width=15
    ws.freeze_panes="D6"
    ws.sheet_properties.pageSetUpPr.fitToPage=True;ws.page_setup.orientation="landscape";ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
    ws.print_area="C2:K60"


def build_checks(wb):
    ws=wb["Checks & Definitions"]
    style_title(ws,"Model checks, metric definitions, and assumption register","Checks are diagnostic outputs only and do not feed assumptions, builds, or dashboard results. All check inputs and rules are synthetic or illustrative.",10)
    section(ws,5,"Integrity checks",8)
    for col,label in enumerate(["Check","Formula result","Status","Reference / interpretation"],start=3):setv(ws,f"{get_column_letter(col)}7",label,kind="label",bold=True,fill=NAVY)
    header(ws,7,3,6)
    checks=[
        (8,"Dashboard EBITDA bridge ties", "='Dashboard'!D31-'Returns'!D20+('Returns'!D20-'Returns'!D20)", "Enterprise value bridge identity is reviewed separately; see line below."),
    ]
    # Actual check list, each calculation independently reads the owned build cells.
    checks=[
        (8,"Entry sponsor equity bridge", "=Returns!D13+Returns!D14-Returns!D15-Returns!D16", "Entry EV plus fees less debt equals sponsor equity."),
        (9,"Exit sponsor proceeds bridge", "=Returns!D20-Returns!D21-Returns!D22-Returns!D23", "Exit EV less exit fees and net debt (gross debt less cash) equals exit equity proceeds."),
        (10,"Dashboard enterprise value bridge", "=Dashboard!D28+Dashboard!D29+Dashboard!D30-Dashboard!D31", "Entry EV plus EBITDA change at entry multiple plus multiple change equals exit EV."),
        (11,"Cash roll-forward, all forecast years", "=MAX(ABS('Cash Flow & Debt'!D18+'Cash Flow & Debt'!D17-'Cash Flow & Debt'!D19-'Cash Flow & Debt'!D21-'Cash Flow & Debt'!D23),ABS('Cash Flow & Debt'!E18+'Cash Flow & Debt'!E17-'Cash Flow & Debt'!E19-'Cash Flow & Debt'!E21-'Cash Flow & Debt'!E23),ABS('Cash Flow & Debt'!F18+'Cash Flow & Debt'!F17-'Cash Flow & Debt'!F19-'Cash Flow & Debt'!F21-'Cash Flow & Debt'!F23),ABS('Cash Flow & Debt'!G18+'Cash Flow & Debt'!G17-'Cash Flow & Debt'!G19-'Cash Flow & Debt'!G21-'Cash Flow & Debt'!G23),ABS('Cash Flow & Debt'!H18+'Cash Flow & Debt'!H17-'Cash Flow & Debt'!H19-'Cash Flow & Debt'!H21-'Cash Flow & Debt'!H23))", "Opening cash plus FCF less amortization and sweep equals ending cash in each year."),
        (12,"Debt roll-forward, all forecast years", "=MAX(ABS('Cash Flow & Debt'!D11-'Cash Flow & Debt'!D19-'Cash Flow & Debt'!D21-'Cash Flow & Debt'!D22),ABS('Cash Flow & Debt'!E11-'Cash Flow & Debt'!E19-'Cash Flow & Debt'!E21-'Cash Flow & Debt'!E22),ABS('Cash Flow & Debt'!F11-'Cash Flow & Debt'!F19-'Cash Flow & Debt'!F21-'Cash Flow & Debt'!F22),ABS('Cash Flow & Debt'!G11-'Cash Flow & Debt'!G19-'Cash Flow & Debt'!G21-'Cash Flow & Debt'!G22),ABS('Cash Flow & Debt'!H11-'Cash Flow & Debt'!H19-'Cash Flow & Debt'!H21-'Cash Flow & Debt'!H22))", "Opening debt less amortization and sweep equals ending debt in each year."),
        (13,"No negative ending gross debt", "=MIN('Cash Flow & Debt'!D22:H22)", "All years must be zero or positive."),
        (14,"No negative ending cash", "=MIN('Cash Flow & Debt'!D23:H23)", "All years must be zero or positive; a negative value signals an unmodeled funding need."),
        (15,"Pricing off removes contribution", "=IF(Controls!D8=\"Off\",SUM('Operating Model'!D17:H17),0)", "If Off, all pricing contributions must be zero."),
        (16,"Commercial off removes contribution", "=IF(Controls!D9=\"Off\",SUM('Operating Model'!D23:H23),0)", "If Off, all commercial contributions must be zero."),
        (17,"AI off removes contribution", "=IF(Controls!D10=\"Off\",SUM('Operating Model'!D32:H32),0)", "If Off, all AI contributions must be zero."),
        (18,"Selected scenario maps to controls", "=IF(OR(Controls!D6=Controls!C15,Controls!D6=Controls!C16,Controls!D6=Controls!C17),0,1)", "Selected case must match exactly one scenario row."),
        (19,"Base case output formulas present", "=COUNT(Returns!D25,Returns!D26)", "MOIC and XIRR are formula outputs."),
    ]
    for row,label,formula,note in checks:
        setv(ws,f"C{row}",label,kind="label")
        setv(ws,f"D{row}",formula,number_format=MONEY)
        if row in (15,16,17,18):
            status_formula=f'=IF(D{row}=0,"PASS","FAIL")'
        elif row==19:
            status_formula=f'=IF(D{row}=Controls!$D$35,"PASS","FAIL")'
        elif row in (13,14):
            status_formula=f'=IF(D{row}>=0,"PASS","FAIL")'
        else:
            status_formula=f'=IF(ABS(D{row})<Controls!$D$34,"PASS","FAIL")'
        setv(ws,f"E{row}",status_formula)
        setv(ws,f"F{row}",note,kind="label")
    ws.conditional_formatting.add("E8:E19",FormulaRule(formula=['E8="PASS"'],fill=PatternFill("solid",fgColor=LIGHT_GREEN)))
    ws.conditional_formatting.add("E8:E19",FormulaRule(formula=['E8="FAIL"'],fill=PatternFill("solid",fgColor=LIGHT_RED)))
    section(ws,22,"KPI formula references",8)
    for col,label in enumerate(["Dashboard KPI","Cell reference","Definition"],start=3):setv(ws,f"{get_column_letter(col)}24",label,kind="label",bold=True,fill=NAVY)
    header(ws,24,3,5)
    refs=[("Entry EBITDA","Dashboard!D8","Synthetic FY2025 EBITDA from Historical & Baseline."),("Entry EV","Dashboard!D9","Entry EBITDA multiplied by the entry multiple."),("Sponsor equity","Dashboard!D10","Entry EV plus fees less opening senior debt."),("Exit EBITDA","Dashboard!D11","Year 5 forecast EBITDA including selected initiatives."),("Exit EV","Dashboard!D12","Exit EBITDA multiplied by scenario exit multiple."),("Exit net debt","Dashboard!D13","Year 5 gross debt less year-end cash."),("MOIC","Dashboard!D14","Exit equity proceeds divided by sponsor equity."),("IRR","Dashboard!D15","XIRR on dated sponsor cash flows."),("Revenue trend","Dashboard!D36:K36","Historical synthetic revenue and forecast operating revenue."),("EBITDA trend","Dashboard!D37:K37","Historical synthetic EBITDA and forecast EBITDA."),("Bridge exit EBITDA","Dashboard!D24","Core EBITDA change plus three net initiative contributions."),("Bridge exit EV","Dashboard!D31","Entry EV plus EBITDA and exit-multiple effects."),("Red RAG count","Dashboard!D44","Count of quarterly initiative rows below visible threshold."),("AI expand gate count","Dashboard!D45","Count of quarterly AI rows meeting adoption and quality threshold logic.")]
    for row,(name,ref,definition) in enumerate(refs,start=25):
        setv(ws,f"C{row}",name,kind="label");setv(ws,f"D{row}",ref,kind="label");setv(ws,f"E{row}",definition,kind="label")
    section(ws,42,"Metric definitions",8)
    definitions=[
        "EBITDA is the synthetic operating earnings proxy before depreciation, interest, and tax. No formal accounting policy is modeled.",
        "Sponsor MOIC is gross exit equity proceeds divided by sponsor equity at entry. No interim distributions are assumed in the seeded case.",
        "Sponsor IRR uses XIRR with actual Excel dates. No fees beyond the displayed entry and exit transaction fees are included.",
        "AI capacity released is addressable annual hours multiplied by adoption ramp. EBITDA benefit is only hours × loaded cost × avoidable share × realization. The residual capacity is not valued.",
        "Synthetic VCP actuals are seeded for demonstration and are not measured results for any real business.",
    ]
    for i,text in enumerate(definitions,start=43):setv(ws,f"C{i}",text,kind="label");ws.merge_cells(start_row=i,start_column=3,end_row=i,end_column=8);ws[f"C{i}"].alignment=Alignment(wrap_text=True,vertical="top");ws.row_dimensions[i].height=30
    section(ws,50,"Synthetic assumption register",8)
    for col,label in enumerate(["Input group","Key assumption","Synthetic rationale"],start=3):setv(ws,f"{get_column_letter(col)}52",label,kind="label",bold=True,fill=NAVY)
    header(ws,52,3,5)
    reg=[("Company baseline","FY2025 revenue and EBITDA","Illustrative scale for a synthetic mid-market services case."),("Operating scenarios","Growth, margin lift, and exit multiple","Synthetic ranges to create meaningfully distinct operating and execution cases."),("Initiative economics","Benefits, ramps, implementation and run cost","Illustrative only; no observed company or project economics."),("Capital structure","Entry multiple, leverage, fees, interest and sweep","Simplified transaction assumptions, not a financing quote."),("AI automation","Addressable hours, cost/hour, avoidable share, realization","Capacity is monetized only when cost avoidance is explicitly assumed."),("Quarterly tracker","Plan, synthetic actual attainment and quality","Seeded for RAG and expand/fix/pause demonstrations."),("Tax and cash conversion","Tax, capex, working capital, depreciation","Simple illustrative rates; not tax or accounting advice.")]
    for row,data in enumerate(reg,start=53):
        for col,val in enumerate(data,start=3):setv(ws,f"{get_column_letter(col)}{row}",val,kind="label")
    for col,width in {"C":47,"D":45,"E":78,"F":80}.items():ws.column_dimensions[col].width=width
    ws.freeze_panes="C8"


def finish_styles(wb):
    thin=Side(style="thin",color=LINE)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None: continue
                cell.alignment=copy(cell.alignment) + Alignment(vertical="center")
                if cell.data_type=="f" and not cell.font.bold:
                    cell.font=Font(name="Arial",size=10,color=GREEN if "!" in cell.value else BLACK)
                elif cell.data_type!="f" and cell.font.name is None:
                    cell.font=Font(name="Arial",size=10,color=BLACK)
                if cell.row>=6 and cell.border==Border():
                    cell.border=Border(bottom=Side(style="hair",color="E6E6E6"))
        ws.sheet_view.zoomScale=90
    # Explicitly mark every formula's font as black or cross-sheet green after all table styling.
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type=="f":
                    color=GREEN if "!" in cell.value else BLACK
                    cell.font=Font(name="Arial",size=cell.font.sz or 10,bold=cell.font.bold,color=color,italic=cell.font.italic)


def build_workbook(output: Path = OUTPUT, *, scenario: str = "Base", toggles: dict[str, str] | None = None):
    output.parent.mkdir(parents=True,exist_ok=True)
    wb=setup_workbook()
    build_start(wb);build_controls(wb);build_history(wb);build_initiatives(wb)
    build_operating_model(wb);build_cashflow(wb);build_returns(wb);build_sensitivity(wb)
    build_vcp(wb);build_dashboard(wb);build_checks(wb);finish_styles(wb)
    # Link the sensitivity grid to the matching AI realization recalculation engine.
    for row in range(23,28):
        for col in range(4,9):
            L=get_column_letter(col); base=36+(col-4)*16
            wb["Sensitivity"][f"{L}{row}"]=f"=(('Operating Model'!$H$35-'Operating Model'!$H$29+'Operating Model'!$H$28*$D${base})*$C{row}*(1-Controls!$D$24)-$I${base+13}+$I${base+14})/'Returns'!$D$16"
            wb["Sensitivity"][f"{L}{row}"].number_format=MULT
            wb["Sensitivity"][f"{L}{row}"].font=Font(name="Arial",size=10,color=GREEN)
    if scenario not in {"Base", "Downside", "Upside"}:
        raise ValueError(f"Unsupported scenario: {scenario}")
    wb["Controls"]["D6"] = scenario
    for key, cell in {"Pricing and mix": "D8", "Commercial efficiency": "D9", "AI finance automation": "D10"}.items():
        if toggles and key in toggles:
            if toggles[key] not in {"On", "Off"}:
                raise ValueError(f"Unsupported toggle for {key}: {toggles[key]}")
            wb["Controls"][cell] = toggles[key]
    wb.save(output)
    return output


if __name__ == "__main__":
    print(build_workbook())
